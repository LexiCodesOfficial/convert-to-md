"""Tests for convert-to-md utility modules and converters."""

from __future__ import annotations

import pathlib
import tempfile

import openpyxl
import pytest
from docx import Document
from pptx import Presentation

from converters.text_converter import TextConverter
from converters.word_converter import WordConverter
from converters.pptx_converter import PptxConverter
from converters.xlsx_converter import XlsxConverter
from utils.file_utils import detect_file_type, get_output_path, collect_files
from utils.markdown_utils import heading, bullet_list, table, clean_text


# ---------------------------------------------------------------------------
# utils.markdown_utils
# ---------------------------------------------------------------------------


class TestMarkdownUtils:
    def test_heading_level1(self):
        assert heading("Hello") == "# Hello"

    def test_heading_level2(self):
        assert heading("World", level=2) == "## World"

    def test_heading_clamps_max(self):
        assert heading("X", level=10) == "###### X"

    def test_heading_clamps_min(self):
        assert heading("X", level=0) == "# X"

    def test_bullet_list(self):
        result = bullet_list(["alpha", "beta"])
        assert result == "- alpha\n- beta"

    def test_bullet_list_skips_empty(self):
        result = bullet_list(["a", "", "b"])
        assert result == "- a\n- b"

    def test_table_structure(self):
        result = table(["H1", "H2"], [["r1c1", "r1c2"]])
        assert "| H1 | H2 |" in result
        assert "| --- | --- |" in result
        assert "| r1c1 | r1c2 |" in result

    def test_table_escapes_pipe(self):
        result = table(["A|B"], [["v|w"]])
        assert "A\\|B" in result
        assert "v\\|w" in result

    def test_table_empty_headers(self):
        assert table([], []) == ""

    def test_clean_text_strips(self):
        assert clean_text("  hello  ") == "hello"

    def test_clean_text_collapses_blank_lines(self):
        assert clean_text("a\n\n\n\nb") == "a\n\nb"


# ---------------------------------------------------------------------------
# utils.file_utils
# ---------------------------------------------------------------------------


class TestFileUtils:
    @pytest.mark.parametrize(
        "filename,expected",
        [
            ("doc.pdf", "pdf"),
            ("doc.docx", "word"),
            ("doc.doc", "word"),
            ("doc.pptx", "pptx"),
            ("doc.ppt", "pptx"),
            ("doc.txt", "text"),
            ("doc.png", "image"),
            ("doc.jpg", "image"),
            ("doc.jpeg", "image"),
            ("doc.mp4", "video"),
            ("doc.xlsx", "xlsx"),
            ("doc.xls", "xlsx"),
            ("doc.unknown", None),
        ],
    )
    def test_detect_file_type(self, filename, expected):
        assert detect_file_type(filename) == expected

    def test_get_output_path(self):
        result = get_output_path("some/dir/report.pdf", "/out")
        assert result == pathlib.Path("/out/report.md")

    def test_get_output_path_with_input_root(self):
        result = get_output_path("/in/sub/report.pdf", "/out", input_root="/in")
        assert result == pathlib.Path("/out/sub/report.md")

    def test_collect_files_single(self, tmp_path):
        f = tmp_path / "notes.txt"
        f.write_text("hello")
        assert collect_files(f) == [f]

    def test_collect_files_dir(self, tmp_path):
        (tmp_path / "a.txt").write_text("a")
        (tmp_path / "b.xlsx").write_text("")
        (tmp_path / "c.unknown").write_text("")
        files = collect_files(tmp_path)
        names = [p.name for p in files]
        assert "a.txt" in names
        assert "c.unknown" not in names

    def test_collect_files_subdir(self, tmp_path):
        sub = tmp_path / "chapter1"
        sub.mkdir()
        (tmp_path / "intro.txt").write_text("intro")
        (sub / "notes.txt").write_text("notes")
        files = collect_files(tmp_path)
        names = [p.name for p in files]
        assert "intro.txt" in names
        assert "notes.txt" in names

    def test_collect_files_nested(self, tmp_path):
        deep = tmp_path / "a" / "b" / "c"
        deep.mkdir(parents=True)
        (deep / "deep.txt").write_text("deep content")
        files = collect_files(tmp_path)
        assert any(p.name == "deep.txt" for p in files)

    def test_collect_files_unsupported_returns_empty(self, tmp_path):
        f = tmp_path / "file.xyz"
        f.write_text("data")
        assert collect_files(f) == []


# ---------------------------------------------------------------------------
# TextConverter
# ---------------------------------------------------------------------------


class TestTextConverter:
    def test_basic_conversion(self, tmp_path):
        src = tmp_path / "notes.txt"
        src.write_text("INTRODUCTION\n\nHello world.\n- item\n")
        out = tmp_path / "notes.md"
        TextConverter().convert(src, out)
        content = out.read_text()
        assert "## Introduction" in content
        assert "Hello world." in content
        assert "- item" in content

    def test_numbered_list(self, tmp_path):
        src = tmp_path / "list.txt"
        src.write_text("1. First\n2. Second\n")
        out = tmp_path / "list.md"
        TextConverter().convert(src, out)
        content = out.read_text()
        assert "1. First" in content
        assert "2. Second" in content

    def test_output_file_created(self, tmp_path):
        src = tmp_path / "simple.txt"
        src.write_text("hello")
        out = tmp_path / "simple.md"
        TextConverter().convert(src, out)
        assert out.exists()


# ---------------------------------------------------------------------------
# WordConverter
# ---------------------------------------------------------------------------


class TestWordConverter:
    def _make_docx(self, path: pathlib.Path) -> None:
        doc = Document()
        doc.add_heading("My Report", level=1)
        doc.add_paragraph("Intro paragraph.")
        doc.add_heading("Section Two", level=2)
        p = doc.add_paragraph("Bullet item", style="List Bullet")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = "Name"
        tbl.rows[0].cells[1].text = "Value"
        tbl.rows[1].cells[0].text = "foo"
        tbl.rows[1].cells[1].text = "bar"
        doc.save(str(path))

    def test_conversion(self, tmp_path):
        src = tmp_path / "report.docx"
        self._make_docx(src)
        out = tmp_path / "report.md"
        WordConverter().convert(src, out)
        content = out.read_text()
        assert "# My Report" in content
        assert "## Section Two" in content
        assert "- Bullet item" in content
        assert "| Name | Value |" in content


# ---------------------------------------------------------------------------
# PptxConverter
# ---------------------------------------------------------------------------


class TestPptxConverter:
    def _make_pptx(self, path: pathlib.Path) -> None:
        prs = Presentation()
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Overview"
        tf = slide.placeholders[1].text_frame
        tf.text = "Point one"
        p = tf.add_paragraph()
        p.text = "Point two"
        prs.save(str(path))

    def test_conversion(self, tmp_path):
        src = tmp_path / "slides.pptx"
        self._make_pptx(src)
        out = tmp_path / "slides.md"
        PptxConverter().convert(src, out)
        content = out.read_text()
        assert "## Slide 1: Overview" in content
        assert "Point one" in content


# ---------------------------------------------------------------------------
# XlsxConverter
# ---------------------------------------------------------------------------


class TestXlsxConverter:
    def _make_xlsx(self, path: pathlib.Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(["Subject", "Score"])
        ws.append(["Math", 90])
        ws.append(["Science", 88])
        wb.save(str(path))

    def test_conversion(self, tmp_path):
        src = tmp_path / "grades.xlsx"
        self._make_xlsx(src)
        out = tmp_path / "grades.md"
        XlsxConverter().convert(src, out)
        content = out.read_text()
        assert "## Sheet: Results" in content
        assert "| Subject | Score |" in content
        assert "Math" in content
        assert "Science" in content

    def test_empty_sheet(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Empty"
        path = tmp_path / "empty.xlsx"
        wb.save(str(path))
        out = tmp_path / "empty.md"
        XlsxConverter().convert(path, out)
        content = out.read_text()
        assert "*Empty sheet.*" in content


# ---------------------------------------------------------------------------
# CLI (main.py)
# ---------------------------------------------------------------------------


class TestCLI:
    def test_help(self):
        import subprocess
        import sys

        result = subprocess.run(
            [sys.executable, "main.py", "--help"], capture_output=True, text=True
        )
        assert result.returncode == 0
        assert "convert-to-md" in result.stdout

    def test_missing_input(self, tmp_path):
        import subprocess
        import sys

        result = subprocess.run(
            [sys.executable, "main.py", str(tmp_path / "nonexistent"), str(tmp_path / "out")],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 1

    def test_batch_txt_xlsx(self, tmp_path):
        import subprocess
        import sys

        inp = tmp_path / "input"
        inp.mkdir()
        out = tmp_path / "output"

        (inp / "notes.txt").write_text("TITLE\n\nContent here.\n")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A", "B"])
        ws.append([1, 2])
        wb.save(str(inp / "data.xlsx"))

        result = subprocess.run(
            [sys.executable, "main.py", str(inp), str(out)],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 0
        assert (out / "notes.md").exists()
        assert (out / "data.md").exists()

    def test_batch_subdir(self, tmp_path):
        import subprocess
        import sys

        inp = tmp_path / "input"
        sub = inp / "chapter1"
        sub.mkdir(parents=True)
        out = tmp_path / "output"

        (inp / "intro.txt").write_text("INTRO\n\nHello.\n")
        (sub / "notes.txt").write_text("NOTES\n\nWorld.\n")

        result = subprocess.run(
            [sys.executable, "main.py", str(inp), str(out)],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 0
        assert (out / "intro.md").exists()
        assert (out / "chapter1" / "notes.md").exists()
